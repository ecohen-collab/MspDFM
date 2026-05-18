"""Génère Audit_Console_Avril_2026.xlsx — détail machine par machine selon Artis vs DATTO."""
import json, os, re
from pathlib import Path
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

TEMP = os.environ.get('TEMP') or '/tmp'
MSP = json.loads(Path(TEMP, 'msp_data.json').read_text(encoding='utf-8'))

recon = {r['code']: r for r in MSP.get('reconciliation', []) if r.get('code')}
datto_by_code = {}
for s in MSP.get('raw_sources', {}).get('datto', []):
    c = str(s.get('artis_code') or '').lstrip('0') or '0'
    if c == '0' or not c: continue
    datto_by_code.setdefault(c, []).append(s)
m365_by_code = {}
for s in MSP.get('raw_sources', {}).get('m365', []):
    c = str(s.get('artis_code') or '').lstrip('0') or '0'
    if c == '0' or not c: continue
    m365_by_code.setdefault(c, []).append(s)

def extract_code(label):
    if not label: return None
    s = str(label)
    m = re.match(r'^\s*0*(\d{3,6})\s*[-\s]', s) or re.search(r'\((\d{3,6})\)', s)
    return m.group(1) if m else None

def artis_billing(code):
    r = recon.get(code) or {}
    a = r.get('artis') or {}
    return {k: int(a.get(k) or 0) for k in ('datto','m365','sophos','veeam_tb','veeam365','vade','soc')}

def fmt_last_seen(ts):
    if not ts: return ''
    try: return datetime.fromtimestamp(float(ts)).strftime('%Y-%m-%d')
    except: return ''

# Excel
src = openpyxl.load_workbook(Path.cwd() / 'Etat console AVRIL 2026.xlsx', data_only=True)
ws_src = src['avril']
sections, current = {}, None
for row in ws_src.iter_rows(values_only=True):
    if all(c is None for c in row): continue
    c0 = row[0]
    if c0 and not row[1] and not row[2] and isinstance(c0, str) and c0.strip().upper() in ('DATTO','M365'):
        current = c0.strip().upper(); sections[current] = []; continue
    if current: sections[current].append(row[:5])

wb = openpyxl.Workbook()
wb.remove(wb.active)

HEADER_FILL = PatternFill('solid', fgColor='1F4E78')
HEADER_FONT = Font(color='FFFFFF', bold=True)
WARN_FILL   = PatternFill('solid', fgColor='FFE699')
OK_FILL     = PatternFill('solid', fgColor='C6EFCE')
BORDER      = Border(left=Side(style='thin', color='CCCCCC'),
                     right=Side(style='thin', color='CCCCCC'),
                     top=Side(style='thin', color='CCCCCC'),
                     bottom=Side(style='thin', color='CCCCCC'))

def add_sheet(name):
    ws = wb.create_sheet(name)
    ws.freeze_panes = 'A2'
    return ws

def write_header(ws, headers):
    for col, h in enumerate(headers, 1):
        c = ws.cell(1, col, h)
        c.fill = HEADER_FILL; c.font = HEADER_FONT
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = BORDER

def autosize(ws):
    for col_letter, max_len in ws.column_dimensions.items(): pass
    for col_cells in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col_cells), default=0)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(max_len+2, 10), 60)

# ── Buckets DATTO ────────────────────────────────────────────────────────────
poste_only, serv_only, mettre, pas_infog = [], [], [], []
for row in sections.get('DATTO', []):
    label, action, comment, qty, _ = (list(row)+[None]*5)[:5]
    code = extract_code(label)
    if not code: continue
    a = str(action or '').lower(); c = str(comment or '').lower().strip()
    rec = (code, str(label).strip(), str(comment or '').strip(), str(action or '').strip())
    if 'à mettre en place' in a or 'a mettre en place' in a: mettre.append(rec)
    elif 'poste uniquement' in c or 'postes uniquement' in c: poste_only.append(rec)
    elif 'serveur uniquement' in c or 'server uniquement' in c or 'servers uniquement' in c: serv_only.append(rec)
    elif "pas d'infogérance" in c or "pas d'infogerance" in c: pas_infog.append(rec)

# ── Sheet DATTO — Postes uniquement (serveurs à résilier sur DATTO) ──────────
ws = add_sheet('DATTO Postes uniquement')
write_header(ws, ['Code', 'Client', 'Artis facturé (postes)', 'DATTO postes', 'DATTO serveurs', 'Statut', 'Serveur DATTO à résilier (hostname)', 'OS', 'Dernière connexion'])
seen = set()
for code, label, _, _ in poste_only:
    if code in seen: continue
    seen.add(code)
    bill = artis_billing(code)
    sites = datto_by_code.get(code, [])
    total_s = sum(int(s.get('servers') or 0) for s in sites)
    total_p = sum(int(s.get('desktops') or 0) + int(s.get('laptops') or 0) for s in sites)
    servers_to_remove = []
    for s in sites:
        for dev in s.get('devices') or []:
            if (dev.get('category') or '').lower() == 'server':
                servers_to_remove.append(dev)
    if not servers_to_remove:
        ws.append([code, label, bill['datto'], total_p, total_s, 'OK (aucun serveur)', '', '', ''])
        ws.cell(ws.max_row, 6).fill = OK_FILL
    else:
        for i, dev in enumerate(servers_to_remove):
            ws.append([
                code if i == 0 else '',
                label if i == 0 else '',
                bill['datto'] if i == 0 else '',
                total_p if i == 0 else '',
                total_s if i == 0 else '',
                f'À résilier {len(servers_to_remove)} serveur(s)' if i == 0 else '',
                dev.get('hostname',''),
                dev.get('os',''),
                fmt_last_seen(dev.get('lastSeen'))
            ])
            if i == 0:
                ws.cell(ws.max_row, 6).fill = WARN_FILL
autosize(ws)

# ── Sheet DATTO — Serveurs uniquement (postes à résilier sur DATTO) ──────────
ws = add_sheet('DATTO Serveurs uniquement')
write_header(ws, ['Code', 'Client', 'Artis facturé (serveurs)', 'DATTO postes', 'DATTO serveurs', 'Statut', 'Poste DATTO à résilier (hostname)', 'OS', 'Dernière connexion'])
seen = set()
for code, label, _, _ in serv_only:
    if code in seen: continue
    seen.add(code)
    bill = artis_billing(code)
    sites = datto_by_code.get(code, [])
    total_s = sum(int(s.get('servers') or 0) for s in sites)
    total_p = sum(int(s.get('desktops') or 0) + int(s.get('laptops') or 0) for s in sites)
    postes_to_remove = []
    for s in sites:
        for dev in s.get('devices') or []:
            if (dev.get('category') or '').lower() in ('laptop','desktop'):
                postes_to_remove.append(dev)
    if not postes_to_remove:
        ws.append([code, label, bill['datto'], total_p, total_s, 'OK (aucun poste)', '', '', ''])
        ws.cell(ws.max_row, 6).fill = OK_FILL
    else:
        for i, dev in enumerate(postes_to_remove):
            ws.append([
                code if i == 0 else '',
                label if i == 0 else '',
                bill['datto'] if i == 0 else '',
                total_p if i == 0 else '',
                total_s if i == 0 else '',
                f'À résilier {len(postes_to_remove)} poste(s)' if i == 0 else '',
                dev.get('hostname',''),
                dev.get('os',''),
                fmt_last_seen(dev.get('lastSeen'))
            ])
            if i == 0:
                ws.cell(ws.max_row, 6).fill = WARN_FILL
autosize(ws)

# ── Sheet DATTO — Pas d'infogérance (TOUTES les machines à résilier) ─────────
ws = add_sheet("DATTO Pas d'infogérance")
write_header(ws, ['Code', 'Client', 'Artis datto facturé', 'Autres services Artis', 'DATTO postes', 'DATTO serveurs', 'Machine à résilier (hostname)', 'Type', 'OS', 'Dernière connexion'])
seen = set()
for code, label, _, _ in pas_infog:
    if code in seen: continue
    seen.add(code)
    bill = artis_billing(code)
    others = ', '.join(f'{k}={v}' for k,v in bill.items() if v > 0 and k != 'datto') or 'aucun'
    sites = datto_by_code.get(code, [])
    total_s = sum(int(s.get('servers') or 0) for s in sites)
    total_p = sum(int(s.get('desktops') or 0) + int(s.get('laptops') or 0) for s in sites)
    all_devs = []
    for s in sites:
        for dev in s.get('devices') or []:
            all_devs.append(dev)
    if not all_devs:
        ws.append([code, label, bill['datto'], others, total_p, total_s, '(aucun appareil)', '', '', ''])
    else:
        for i, dev in enumerate(all_devs):
            ws.append([
                code if i == 0 else '',
                label if i == 0 else '',
                bill['datto'] if i == 0 else '',
                others if i == 0 else '',
                total_p if i == 0 else '',
                total_s if i == 0 else '',
                dev.get('hostname',''),
                dev.get('category',''),
                dev.get('os',''),
                fmt_last_seen(dev.get('lastSeen'))
            ])
autosize(ws)

# ── Sheet DATTO — À mettre en place ──────────────────────────────────────────
ws = add_sheet('DATTO À mettre en place')
write_header(ws, ['Code', 'Client', 'Commentaire facturation', 'Artis datto facturé', 'DATTO postes actuels', 'DATTO serveurs actuels', 'Machine déjà présente (hostname)', 'Type', 'OS', 'Dernière connexion'])
for code, label, comment, _ in mettre:
    bill = artis_billing(code)
    sites = datto_by_code.get(code, [])
    total_s = sum(int(s.get('servers') or 0) for s in sites)
    total_p = sum(int(s.get('desktops') or 0) + int(s.get('laptops') or 0) for s in sites)
    all_devs = [dev for s in sites for dev in (s.get('devices') or [])]
    if not all_devs:
        ws.append([code, label, comment, bill['datto'], 0, 0, '(rien sur DATTO)', '', '', ''])
    else:
        for i, dev in enumerate(all_devs):
            ws.append([
                code if i == 0 else '',
                label if i == 0 else '',
                comment if i == 0 else '',
                bill['datto'] if i == 0 else '',
                total_p if i == 0 else '',
                total_s if i == 0 else '',
                dev.get('hostname',''),
                dev.get('category',''),
                dev.get('os',''),
                fmt_last_seen(dev.get('lastSeen'))
            ])
autosize(ws)

# ── M365 — À résilier ────────────────────────────────────────────────────────
m365_resilier, m365_mettre, m365_voir = [], [], []
for row in sections.get('M365', []):
    label, action, comment, qty, _ = (list(row)+[None]*5)[:5]
    code = extract_code(label)
    if not code: continue
    a = str(action or '').lower()
    rec = (code, str(label).strip(), str(comment or '').strip(), str(action or '').strip())
    if 'à résilier' in a or 'a resilier' in a: m365_resilier.append(rec)
    elif 'à mettre en place' in a or 'a mettre en place' in a: m365_mettre.append(rec)
    elif 'à voir' in a or 'a voir' in a: m365_voir.append(rec)

ws = add_sheet('M365 À résilier')
write_header(ws, ['Code', 'Client', 'Commentaire facturation', 'Artis m365 facturé', 'SKU TDSynnex actif', 'Licences', 'Date de fin', 'Status'])
seen = set()
for code, label, comment, _ in m365_resilier:
    if code in seen: continue
    seen.add(code)
    bill = artis_billing(code)
    subs = m365_by_code.get(code, [])
    if not subs:
        ws.append([code, label, comment, bill['m365'], '(aucune sub TDSynnex)', 0, '', ''])
        ws.cell(ws.max_row, 5).fill = OK_FILL
    else:
        for i, s in enumerate(subs):
            ws.append([
                code if i == 0 else '',
                label if i == 0 else '',
                comment if i == 0 else '',
                bill['m365'] if i == 0 else '',
                s.get('subscription',''),
                s.get('licenses',0),
                s.get('end_date',''),
                s.get('status','')
            ])
            ws.cell(ws.max_row, 5).fill = WARN_FILL
autosize(ws)

# ── M365 — À mettre en place ─────────────────────────────────────────────────
ws = add_sheet('M365 À mettre en place')
write_header(ws, ['Code', 'Client', 'Commentaire (vendu)', 'Artis m365 facturé', 'TDSynnex actuel (SKU)', 'TDSynnex actuel (licences)', 'Verdict'])
seen = set()
for code, label, comment, _ in m365_mettre:
    if code in seen: continue
    seen.add(code)
    bill = artis_billing(code)
    subs = m365_by_code.get(code, [])
    if not subs:
        ws.append([code, label, comment, bill['m365'], '(rien)', 0, 'À créer from scratch'])
        ws.cell(ws.max_row, 7).fill = WARN_FILL
    else:
        for i, s in enumerate(subs):
            verdict = 'À compléter' if i == 0 else ''
            ws.append([
                code if i == 0 else '',
                label if i == 0 else '',
                comment if i == 0 else '',
                bill['m365'] if i == 0 else '',
                s.get('subscription',''),
                s.get('licenses',0),
                verdict
            ])
            if i == 0: ws.cell(ws.max_row, 7).fill = OK_FILL
autosize(ws)

# ── M365 — À voir ────────────────────────────────────────────────────────────
ws = add_sheet('M365 À voir')
write_header(ws, ['Code', 'Client', 'Commentaire', 'Artis m365 facturé', 'TDSynnex actuel (SKU)', 'TDSynnex actuel (licences)'])
seen = set()
for code, label, comment, _ in m365_voir:
    if code in seen: continue
    seen.add(code)
    bill = artis_billing(code)
    subs = m365_by_code.get(code, [])
    if not subs:
        ws.append([code, label, comment, bill['m365'], '(rien)', 0])
    else:
        for i, s in enumerate(subs):
            ws.append([
                code if i == 0 else '',
                label if i == 0 else '',
                comment if i == 0 else '',
                bill['m365'] if i == 0 else '',
                s.get('subscription',''),
                s.get('licenses',0)
            ])
autosize(ws)

# ── Sheet résumé ─────────────────────────────────────────────────────────────
ws = wb.create_sheet('Résumé', 0)
ws.append(['Audit Console Avril 2026 — généré le ' + datetime.now().strftime('%Y-%m-%d %H:%M')])
ws.append([''])
ws.append(['Source données : msp_data.json (' + str(MSP.get('updated_at_display','?')) + ')'])
ws.append([''])
ws.append(['Onglet', 'Nb clients', 'Description'])
ws.append(['DATTO Postes uniquement', len(set(c for c,_,_,_ in poste_only)), 'Clients facturés en postes seuls sur Artis — détail des serveurs DATTO à résilier'])
ws.append(['DATTO Serveurs uniquement', len(set(c for c,_,_,_ in serv_only)), 'Clients facturés en serveurs seuls sur Artis — détail des postes DATTO à résilier'])
ws.append(["DATTO Pas d'infogérance", len(set(c for c,_,_,_ in pas_infog)), "Clients sans infogérance Artis — toutes les machines DATTO à résilier"])
ws.append(['DATTO À mettre en place', len(set(c for c,_,_,_ in mettre)), "Clients vendus mais pas (ou partiellement) déployés DATTO"])
ws.append(['M365 À résilier', len(set(c for c,_,_,_ in m365_resilier)), 'Licences TDSynnex actives à arrêter'])
ws.append(['M365 À mettre en place', len(set(c for c,_,_,_ in m365_mettre)), "Vendu Artis vs déployé TDSynnex"])
ws.append(['M365 À voir', len(set(c for c,_,_,_ in m365_voir)), 'Anomalies/écarts ponctuels'])
for col in range(1,4): ws.column_dimensions[chr(64+col)].width = [35, 12, 80][col-1]
for r in range(5, 13): ws.cell(r, 1).font = Font(bold=(r==5))

# Save
out = Path.cwd() / 'Audit_Console_Avril_2026.xlsx'
wb.save(out)
print(f'OK : {out} ({out.stat().st_size} bytes)')
